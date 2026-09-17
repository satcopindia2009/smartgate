"""Student + authorized-pickup CSV/Excel import (locked Viren/Admin template)."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Any, Optional

from app.config import (
    PICKUP_CONSENT_VERSION,
    PRANAY_SCHOOL_CODE,
    PRANAY_SCHOOL_ID,
    RESERVED_SCHOOL_ID,
    WATERMARK,
)
from app.errors import AppError
from app.models import CustodyFlag, IdType, PickupRelation
from app.pickup_match import default_custody_flag, denormalize_blocked_by_custody
from app.util import last4, normalize_mobile, now_iso
from app import store

# Pickup/PM SoT — one row = one authorized person; student fields repeat.
LOCKED_HEADERS = (
    "school_code",
    "student_external_id",
    "student_name",
    "class",
    "section",
    "person_name",
    "relation",
    "mobile",
    "id_last4",
    "id_type",
    "effective_from",
    "effective_to",
    "custody_flag",
    "gate_instruction",
    "allowed_person_mobiles",
    "blocked_person_mobiles",
    "consent_version",
    "consent_at",
    "person_active",
    "legal_hold",
)
CSV_CONTRACT = (
    "Locked CSV/Excel headers (snake_case, one row = one authorized person): "
    + ", ".join(LOCKED_HEADERS)
    + ". school_code maps to schoolId (JWT is source of truth). "
    "student_external_id→studentId, student_name→name, person_name→pickup name, "
    "consent_version/consent_at→pickupConsentVersion/At. "
    "relation: parent|guardian|sibling|relative|other. "
    "Upsert student by (schoolId + student_external_id) and person by "
    "(schoolId + student_external_id + mobile). "
    "mode=validate is dry-run; mode=commit writes."
)

# Accept aliases → locked names. Do not document a second thinner template.
_ALIASES = {
    "schoolid": "school_code",
    "school_id": "school_code",
    "schoolcode": "school_code",
    "studentid": "student_external_id",
    "student_id": "student_external_id",
    "studentexternalid": "student_external_id",
    "name": "student_name",
    "studentname": "student_name",
    "personname": "person_name",
    "idlast4": "id_last4",
    "idtype": "id_type",
    "effectivefrom": "effective_from",
    "effectiveto": "effective_to",
    "custodyflag": "custody_flag",
    "gateinstruction": "gate_instruction",
    "allowedpersonmobiles": "allowed_person_mobiles",
    "blockedpersonmobiles": "blocked_person_mobiles",
    "pickupconsentversion": "consent_version",
    "consentversion": "consent_version",
    "pickupconsentat": "consent_at",
    "consentat": "consent_at",
    "active": "person_active",
    "personactive": "person_active",
    "legalhold": "legal_hold",
}

COURT_DOC_COLUMNS = {
    "court_doc",
    "court_pdf",
    "court_order_doc",
    "court_order_document",
    "court_document",
    "courtdoc",
    "court_order_pdf",
    "court_narrative",
    "case_pdf",
    "order_pdf",
    "custody_document",
    "courtdocument",
    "court_order_file",
    "narrative",
}

RELATION_VALUES = {r.value for r in PickupRelation}
ID_TYPE_VALUES = {i.value for i in IdType}
CUSTODY_VALUES = {c.value for c in CustodyFlag}
STUDENT_KEYS = (
    "school_code",
    "student_external_id",
    "student_name",
    "class",
    "section",
    "legal_hold",
)
PERSON_KEYS = (
    "person_name",
    "relation",
    "mobile",
    "id_last4",
    "id_type",
    "effective_from",
    "effective_to",
    "consent_version",
    "consent_at",
    "person_active",
)


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0 and value.microsecond == 0:
            return value.date().isoformat()
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(name: Any) -> str:
    raw = _cell_str(name).lstrip("\ufeff").strip()
    key = re.sub(r"[^a-z0-9]+", "", raw.lower())
    if key in COURT_DOC_COLUMNS or raw.lower().replace(" ", "_") in COURT_DOC_COLUMNS:
        raise AppError(
            "VALIDATION",
            f"Court-document column '{raw}' is rejected — do not import court PDFs/narratives",
            400,
            {"column": raw},
        )
    locked = _ALIASES.get(key, raw.strip().lower().replace(" ", "_"))
    return locked


def _parse_bool(raw: str, default: bool, field: str) -> bool:
    if raw is None or str(raw).strip() == "":
        return default
    s = str(raw).strip().lower()
    if s in ("1", "true", "yes", "y", "active"):
        return True
    if s in ("0", "false", "no", "n", "inactive"):
        return False
    raise ValueError(f"{field} must be true/false")


def _filename_kind(filename: Optional[str], content_type: Optional[str]) -> str:
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm") or "spreadsheet" in ctype:
        return "xlsx"
    if name.endswith(".csv") or "csv" in ctype or ctype in (
        "text/plain",
        "application/octet-stream",
        "",
    ):
        return "csv"
    raise AppError(
        "VALIDATION",
        "Roster import accepts .csv or .xlsx only",
        400,
        {"filename": filename, "contentType": content_type},
    )


def _read_tabular(data: bytes, filename: Optional[str], content_type: Optional[str]) -> list[dict[str, str]]:
    if not data or not data.strip():
        raise AppError("VALIDATION", "Import file is empty", 400, {"filename": filename})
    kind = _filename_kind(filename, content_type)
    if kind == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise AppError("VALIDATION", "openpyxl is required for .xlsx import", 400) from e
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise AppError("VALIDATION", f"Could not read Excel file: {e}", 400) from e
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as e:
            raise AppError("VALIDATION", "Excel sheet is empty", 400) from e
        headers = [_norm_header(h) for h in header_row]
        if not any(headers):
            raise AppError("VALIDATION", "Excel header row is empty", 400)
        out: list[dict[str, str]] = []
        for excel_row in rows_iter:
            mapped = {
                headers[i]: _cell_str(excel_row[i] if i < len(excel_row) else "")
                for i in range(len(headers))
                if headers[i]
            }
            mapped["_row"] = str(len(out) + 2)
            out.append(mapped)
        return out

    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise AppError("VALIDATION", "CSV header row is missing", 400)
    reader.fieldnames = [_norm_header(h) for h in reader.fieldnames]
    out = []
    for i, row in enumerate(reader, start=2):
        mapped = {_norm_header(k): _cell_str(v) for k, v in (row or {}).items() if k}
        mapped["_row"] = str(i)
        out.append(mapped)
    return out


def _blank_row(row: dict[str, str]) -> bool:
    return not any((row.get(k) or "").strip() for k in LOCKED_HEADERS)


def _err(
    row_no: int,
    message: str,
    field: Optional[str] = None,
    file: Optional[str] = None,
    code: str = "VALIDATION",
) -> dict:
    item: dict[str, Any] = {"row": row_no, "message": message, "code": code}
    if field:
        item["field"] = field
    if file:
        item["file"] = file
    return item


def _check_school_code(row_code: str, jwt_school: str) -> Optional[str]:
    raw = (row_code or "").strip()
    if not raw:
        return None
    if raw.upper() in {"SCH-PRANAY-PUNE-01", "PRANAY-PUNE"}:
        return "SCH-PRANAY-PUNE-01 is not a valid school_code — use PRANAY / SCH-PRANAY-01"
    found = store.get_school_by_code(raw)
    if jwt_school == PRANAY_SCHOOL_ID:
        if raw.upper() not in {PRANAY_SCHOOL_CODE, PRANAY_SCHOOL_ID}:
            return (
                f"school_code must be {PRANAY_SCHOOL_CODE} for tenant {PRANAY_SCHOOL_ID}"
            )
        return None
    if found:
        if found["id"] != jwt_school:
            return (
                f"school_code '{raw}' is {found['id']}, not JWT schoolId '{jwt_school}' "
                "(JWT is source of truth)"
            )
        return None
    if raw != jwt_school and raw.upper() != jwt_school.upper():
        return f"school_code '{raw}' does not match JWT schoolId '{jwt_school}' (JWT is source of truth)"
    return None


def _split_mobiles(raw: str) -> list[str]:
    out: list[str] = []
    for part in re.split(r"[,;|/]+", raw or ""):
        token = part.strip()
        if not token:
            continue
        mobile = normalize_mobile(token)
        if not mobile:
            raise ValueError(f"Invalid mobile '{token}' — expected IN 10-digit or E.164")
        if mobile not in out:
            out.append(mobile)
    return out


def import_locked_rows(
    data: bytes,
    *,
    school_id: str,
    user: dict,
    filename: Optional[str] = None,
    content_type: Optional[str] = None,
    file_label: str = "file",
    mode: str = "commit",
    require_person: bool = False,
) -> dict:
    commit = mode == "commit"
    rows = _read_tabular(data, filename, content_type)
    headers = {k for r in rows for k in r.keys() if k != "_row"}
    if "student_external_id" not in headers:
        raise AppError(
            "VALIDATION",
            "Missing required column student_external_id (locked import template)",
            400,
            {"contract": CSV_CONTRACT, "filename": filename},
        )
    created = 0
    updated = 0
    errors: list[dict] = []
    ts = now_iso()
    pending_custody: dict[str, dict] = {}
    seen_students: set[str] = set()
    role = user.get("role")
    user_id = user.get("id") or "unknown"

    for row in rows:
        row_no = int(row.get("_row") or 0)
        if _blank_row(row):
            continue
        mismatch = _check_school_code(row.get("school_code", ""), school_id)
        if mismatch:
            errors.append(_err(row_no, mismatch, "school_code", file_label))
            continue
        roster_id = (row.get("student_external_id") or "").strip()
        student_name = (row.get("student_name") or "").strip()
        klass = (row.get("class") or "").strip()
        section = (row.get("section") or "").strip()
        if not roster_id:
            errors.append(_err(row_no, "student_external_id is required", "student_external_id", file_label))
            continue
        has_student = bool(student_name or klass or section or (row.get("legal_hold") or "").strip())
        has_person = any((row.get(k) or "").strip() for k in PERSON_KEYS)
        if require_person:
            has_person = True
        if roster_id not in seen_students and (
            has_student or not store.get_student_by_roster_id(school_id, roster_id)
        ):
            seen_students.add(roster_id)
            if not student_name:
                errors.append(_err(row_no, "student_name is required", "student_name", file_label))
                continue
            if not klass:
                errors.append(_err(row_no, "class is required", "class", file_label))
                continue
            if not section:
                errors.append(_err(row_no, "section is required", "section", file_label))
                continue
            try:
                legal_hold = _parse_bool(row.get("legal_hold", ""), False, "legal_hold")
            except ValueError as e:
                errors.append(_err(row_no, str(e), "legal_hold", file_label))
                continue
            existing = store.get_student_by_roster_id(school_id, roster_id)
            if existing:
                updated += 1
                if commit:
                    existing.update(
                        {
                            "name": student_name,
                            "class": klass,
                            "section": section,
                            "legalHold": legal_hold,
                            "updatedAt": ts,
                        }
                    )
                    store.put_student(existing)
            else:
                created += 1
                if commit:
                    sid = f"STU-{store.next_seq('student_seq'):04d}"
                    store.put_student(
                        {
                            "id": sid,
                            "schoolId": school_id,
                            "studentId": roster_id,
                            "name": student_name,
                            "class": klass,
                            "section": section,
                            "active": True,
                            "enrollmentEndedAt": None,
                            "legalHold": legal_hold,
                            "createdAt": ts,
                            "updatedAt": ts,
                            "importedByUserId": user_id,
                        }
                    )
                    store.put_custody_flag(default_custody_flag(sid, school_id))

        student = store.get_student_by_roster_id(school_id, roster_id)
        if not student and not commit:
            # dry-run: synthesize a handle for later person/custody checks
            student = {
                "id": f"DRY-{roster_id}",
                "schoolId": school_id,
                "studentId": roster_id,
            }

        if has_person:
            person_name = (row.get("person_name") or "").strip()
            relation = (row.get("relation") or "").strip().lower()
            mobile_raw = (row.get("mobile") or "").strip()
            consent_ver = (row.get("consent_version") or "").strip()
            consent_at = (row.get("consent_at") or "").strip()
            if not person_name:
                errors.append(_err(row_no, "person_name is required", "person_name", file_label))
                continue
            if relation not in RELATION_VALUES:
                errors.append(
                    _err(
                        row_no,
                        f"relation must be one of {sorted(RELATION_VALUES)}",
                        "relation",
                        file_label,
                    )
                )
                continue
            mobile = normalize_mobile(mobile_raw)
            if not mobile:
                errors.append(
                    _err(row_no, "Invalid mobile — expected IN 10-digit or E.164", "mobile", file_label)
                )
                continue
            if not consent_ver:
                errors.append(_err(row_no, "consent_version is required", "consent_version", file_label))
                continue
            if not consent_at:
                errors.append(_err(row_no, "consent_at is required", "consent_at", file_label))
                continue
            try:
                person_active = _parse_bool(row.get("person_active", ""), True, "person_active")
            except ValueError as e:
                errors.append(_err(row_no, str(e), "person_active", file_label))
                continue
            id_type_raw = (row.get("id_type") or "").strip()
            id_type = None
            if id_type_raw:
                match = next((v for v in ID_TYPE_VALUES if v.lower() == id_type_raw.lower()), None)
                if not match:
                    errors.append(
                        _err(
                            row_no,
                            f"id_type must be one of {sorted(ID_TYPE_VALUES)}",
                            "id_type",
                            file_label,
                        )
                    )
                    continue
                id_type = match
            if not student:
                errors.append(
                    _err(
                        row_no,
                        f"Unknown student_external_id '{roster_id}' for this school",
                        "student_external_id",
                        file_label,
                    )
                )
                continue
            existing_p = (
                store.get_authorized_person_by_mobile(school_id, student["id"], mobile)
                if not student["id"].startswith("DRY-")
                else None
            )
            payload = {
                "schoolId": school_id,
                "studentId": student["id"],
                "name": person_name,
                "relation": relation,
                "mobile": mobile,
                "idType": id_type,
                "idNumber": None,
                "idLast4": (row.get("id_last4") or "").strip() or None,
                "active": person_active,
                "effectiveFrom": (row.get("effective_from") or "").strip() or None,
                "effectiveTo": (row.get("effective_to") or "").strip() or None,
                "pickupConsentVersion": consent_ver,
                "pickupConsentAt": consent_at,
                "updatedByUserId": user_id,
                "updatedAt": ts,
            }
            if existing_p:
                updated += 1
                if commit:
                    existing_p.update(payload)
                    store.put_authorized_person(existing_p)
                    denormalize_blocked_by_custody(student["id"], school_id)
            else:
                created += 1
                if commit:
                    pid = f"APP-{store.next_seq('person_seq'):04d}"
                    store.put_authorized_person(
                        {
                            "id": pid,
                            "photoRef": None,
                            "blockedByCustody": False,
                            "createdByUserId": user_id,
                            "createdAt": ts,
                            **payload,
                        }
                    )
                    denormalize_blocked_by_custody(student["id"], school_id)

        flag_raw = (row.get("custody_flag") or "").strip().lower()
        if flag_raw or (row.get("gate_instruction") or "").strip() or (
            row.get("allowed_person_mobiles") or ""
        ).strip() or (row.get("blocked_person_mobiles") or "").strip():
            if flag_raw and flag_raw not in CUSTODY_VALUES:
                errors.append(
                    _err(
                        row_no,
                        f"custody_flag must be one of {sorted(CUSTODY_VALUES)}",
                        "custody_flag",
                        file_label,
                    )
                )
                continue
            pending_custody[roster_id] = {
                "row": row_no,
                "flag": flag_raw or "none",
                "gate_instruction": (row.get("gate_instruction") or "").strip(),
                "allowed": row.get("allowed_person_mobiles") or "",
                "blocked": row.get("blocked_person_mobiles") or "",
                "file": file_label,
            }

    for roster_id, spec in pending_custody.items():
        student = store.get_student_by_roster_id(school_id, roster_id)
        if not student:
            if commit:
                errors.append(
                    _err(
                        spec["row"],
                        f"Unknown student_external_id '{roster_id}' for custody_flag",
                        "student_external_id",
                        spec["file"],
                    )
                )
            continue
        flag = spec["flag"]
        instruction = spec["gate_instruction"]
        if flag == "court_order" and not instruction:
            errors.append(
                _err(
                    spec["row"],
                    "court_order requires a non-empty gate_instruction (F6 fail-closed)",
                    "gate_instruction",
                    spec["file"],
                )
            )
            continue
        if flag == "court_order" and role != "security_head":
            errors.append(
                _err(
                    spec["row"],
                    "court_order custody flag write requires security_head",
                    "custody_flag",
                    spec["file"],
                )
            )
            continue
        try:
            allowed_m = _split_mobiles(spec["allowed"])
            blocked_m = _split_mobiles(spec["blocked"])
        except ValueError as e:
            errors.append(_err(spec["row"], str(e), file=spec["file"]))
            continue
        people = store.list_authorized_people(school_id, student["id"])
        by_mobile = {normalize_mobile(p.get("mobile")): p["id"] for p in people}

        def _resolve(mobiles: list[str], field: str) -> Optional[list[str]]:
            ids = []
            for m in mobiles:
                pid = by_mobile.get(m)
                if not pid:
                    errors.append(
                        _err(
                            spec["row"],
                            f"{field} mobile {m} is not an authorized person on this student",
                            field,
                            spec["file"],
                        )
                    )
                    return None
                ids.append(pid)
            return ids

        allowed_ids = _resolve(allowed_m, "allowed_person_mobiles") if allowed_m else None
        blocked_ids = _resolve(blocked_m, "blocked_person_mobiles")
        if allowed_m and allowed_ids is None:
            continue
        if blocked_m and blocked_ids is None:
            continue
        if commit:
            store.put_custody_flag(
                {
                    "studentId": student["id"],
                    "schoolId": school_id,
                    "flag": flag,
                    "gateInstruction": instruction,
                    "blockedPersonIds": blocked_ids or [],
                    "allowedPersonIds": allowed_ids,
                    "updatedByUserId": user_id,
                    "updatedAt": ts,
                }
            )
            denormalize_blocked_by_custody(student["id"], school_id)
            updated += 1

    return {"created": created, "updated": updated, "errors": errors}


def import_students(
    data: bytes,
    *,
    school_id: str,
    user_id: str,
    filename: Optional[str] = None,
    content_type: Optional[str] = None,
    file_label: str = "students",
    user: Optional[dict] = None,
    mode: str = "commit",
) -> dict:
    actor = user or {"id": user_id, "role": "admin"}
    return import_locked_rows(
        data,
        school_id=school_id,
        user=actor,
        filename=filename,
        content_type=content_type,
        file_label=file_label,
        mode=mode,
        require_person=False,
    )


def import_pickup(
    data: bytes,
    *,
    school_id: str,
    user_id: str,
    filename: Optional[str] = None,
    content_type: Optional[str] = None,
    file_label: str = "pickup",
    user: Optional[dict] = None,
    mode: str = "commit",
) -> dict:
    actor = user or {"id": user_id, "role": "admin"}
    return import_locked_rows(
        data,
        school_id=school_id,
        user=actor,
        filename=filename,
        content_type=content_type,
        file_label=file_label,
        mode=mode,
        require_person=True,
    )


def merge_import_results(
    students: Optional[dict] = None,
    pickup: Optional[dict] = None,
    *,
    unified: Optional[dict] = None,
    mode: str = "commit",
    user: Optional[dict] = None,
    filenames: Optional[list[str]] = None,
    school_id: Optional[str] = None,
) -> dict:
    parts = [p for p in (students, pickup, unified) if p]
    created = sum(int(p.get("created") or 0) for p in parts)
    updated = sum(int(p.get("updated") or 0) for p in parts)
    errors: list[dict] = []
    for p in parts:
        errors.extend(p.get("errors") or [])
    audit = {
        "importedByUserId": (user or {}).get("id"),
        "importedAt": now_iso(),
        "filenames": [f for f in (filenames or []) if f],
        "mode": mode,
        "created": created,
        "updated": updated,
        "errorCount": len(errors),
        "schoolId": school_id,
    }
    if school_id:
        stored = store.add_roster_import(audit)
        audit["id"] = stored["id"]
    school = store.get_school(school_id) if school_id else None
    out = {
        "created": created,
        "imported": created,
        "updated": updated,
        "failed": len(errors),
        "errors": errors,
        "schoolId": school_id,
        "school_code": (school or {}).get("schoolCode") if school else None,
        "mode": mode,
        "dryRun": mode != "commit",
        "audit": audit,
        "meta": {
            "watermark": WATERMARK,
            "reservedSchoolId": RESERVED_SCHOOL_ID,
            "pranaySchoolId": PRANAY_SCHOOL_ID,
            "pranaySchoolCode": PRANAY_SCHOOL_CODE,
            "csvContract": CSV_CONTRACT,
            "defaultConsentVersion": PICKUP_CONSENT_VERSION,
            "lockedHeaders": list(LOCKED_HEADERS),
        },
    }
    if students is not None:
        out["students"] = students
    if pickup is not None:
        out["pickup"] = pickup
    return out


async def read_upload(upload) -> tuple[bytes, Optional[str], Optional[str]]:
    data = await upload.read()
    return data, getattr(upload, "filename", None), getattr(upload, "content_type", None)
