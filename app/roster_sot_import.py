"""Pickup SoT one-file roster import (school_code + person rows)."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any, Optional

from app.config import PICKUP_CONSENT_VERSION, WATERMARK
from app.errors import AppError
from app.models import IdType, PickupRelation
from app.pickup_match import default_custody_flag, denormalize_blocked_by_custody
from app.util import last4, normalize_mobile, now_iso
from app import store

# Product lock: CSV school_code -> schoolId
SCHOOL_CODE_MAP = {
    "PRANAY": "SCH-PRANAY-01",
}
RESERVED_DEMO = "SCH-DEMO-01"
DEMO_CODES = {"DEMO", "SCH-DEMO-01", "SCH_DEMO_01"}

COLUMN_ALIASES = {
    "schoolCode": "school_code",
    "school_code": "school_code",
    "studentExternalId": "student_external_id",
    "student_external_id": "student_external_id",
    "studentId": "student_external_id",
    "student_name": "student_name",
    "studentName": "student_name",
    "class": "class",
    "section": "section",
    "person_name": "person_name",
    "personName": "person_name",
    "relation": "relation",
    "mobile": "mobile",
    "id_last4": "id_last4",
    "idLast4": "id_last4",
    "id_type": "id_type",
    "idType": "id_type",
    "effective_from": "effective_from",
    "effectiveFrom": "effective_from",
    "effective_to": "effective_to",
    "effectiveTo": "effective_to",
    "custody_flag": "custody_flag",
    "custodyFlag": "custody_flag",
    "custody": "custody_flag",
    "gate_instruction": "gate_instruction",
    "gateInstruction": "gate_instruction",
    "allowed_person_mobiles": "allowed_person_mobiles",
    "allowedPersonMobiles": "allowed_person_mobiles",
    "blocked_person_mobiles": "blocked_person_mobiles",
    "blockedPersonMobiles": "blocked_person_mobiles",
    "consent_version": "consent_version",
    "consentVersion": "consent_version",
    "pickupConsentVersion": "consent_version",
    "consent_at": "consent_at",
    "consentAt": "consent_at",
    "pickupConsentAt": "consent_at",
    "person_active": "person_active",
    "personActive": "person_active",
    "legal_hold": "legal_hold",
    "legalHold": "legal_hold",
}

SOT_REQUIRED = (
    "school_code",
    "student_external_id",
    "student_name",
    "class",
    "section",
    "person_name",
    "relation",
    "mobile",
    "consent_version",
    "consent_at",
)
SOT_OPTIONAL = (
    "id_last4",
    "id_type",
    "effective_from",
    "effective_to",
    "custody_flag",
    "gate_instruction",
    "allowed_person_mobiles",
    "blocked_person_mobiles",
    "person_active",
    "legal_hold",
)
FORBIDDEN_COLS = {
    "court_pdf",
    "case_narrative",
    "document_url",
    "opposing_party_address",
}
RELATION_VALUES = {r.value for r in PickupRelation}
ID_TYPE_VALUES = {i.value for i in IdType}
CUSTODY_VALUES = {"none", "restricted", "court_order"}


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.date().isoformat()
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(name: Any) -> str:
    return _cell(name).lstrip("\ufeff").strip()


def _parse_bool(raw: str, default: bool, field: str) -> bool:
    if raw is None or str(raw).strip() == "":
        return default
    s = str(raw).strip().lower()
    if s in ("1", "true", "yes", "y"):
        return True
    if s in ("0", "false", "no", "n"):
        return False
    raise ValueError(f"{field} must be Y/N")


def _read_rows(data: bytes, filename: Optional[str], content_type: Optional[str]) -> list[dict[str, str]]:
    if not data or not data.strip():
        raise AppError("VALIDATION", "Import file is empty", 400)
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    if name.endswith(".xlsx") or "spreadsheet" in ctype:
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise AppError("VALIDATION", "openpyxl is required for .xlsx import", 400) from e
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration as e:
            raise AppError("VALIDATION", "Excel sheet is empty", 400) from e
        headers = [_norm_header(h) for h in header]
        out = []
        for excel_row in it:
            raw = {
                headers[i]: _cell(excel_row[i] if i < len(excel_row) else "")
                for i in range(len(headers))
                if headers[i]
            }
            mapped = {}
            for k, v in raw.items():
                canon = COLUMN_ALIASES.get(k, k)
                if canon not in mapped or k == canon:
                    mapped[canon] = v
            mapped["_row"] = str(len(out) + 2)
            out.append(mapped)
        return out
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise AppError("VALIDATION", "CSV header row is missing", 400)
    reader.fieldnames = [_norm_header(h) for h in reader.fieldnames]
    out = []
    for i, row in enumerate(reader, start=2):
        raw = {_norm_header(k): _cell(v) for k, v in (row or {}).items() if k}
        mapped = {}
        for k, v in raw.items():
            canon = COLUMN_ALIASES.get(k, k)
            if canon not in mapped or k == canon:
                mapped[canon] = v
        mapped["_row"] = str(i)
        out.append(mapped)
    return out


def _err(row: int, message: str, field: Optional[str] = None, code: str = "VALIDATION") -> dict:
    item = {"row": row, "field": field, "code": code, "message": message}
    return {k: v for k, v in item.items() if v is not None}


def resolve_school(*, jwt_school_id: str, rows: list[dict[str, str]]) -> str:
    codes = set()
    for r in rows:
        c = (r.get("school_code") or "").strip().upper()
        if c:
            codes.add(c)
    if not codes:
        if jwt_school_id == RESERVED_DEMO:
            raise AppError(
                "FORBIDDEN",
                "Roster import into SCH-DEMO-01 is blocked on this path",
                403,
                {"jwt": jwt_school_id},
            )
        return jwt_school_id

    if len(codes) > 1:
        raise AppError(
            "VALIDATION",
            f"Mixed school_code values in file: {', '.join(sorted(codes))}",
            400,
            {"school_codes": sorted(codes)},
        )
    code = next(iter(codes))

    if code in DEMO_CODES:
        raise AppError(
            "FORBIDDEN",
            "school_code DEMO / SCH-DEMO-01 is not allowed for this import path",
            403,
            {"school_code": code, "jwt": jwt_school_id, "allowed": list(SCHOOL_CODE_MAP)},
        )

    mapped = SCHOOL_CODE_MAP.get(code)
    if not mapped:
        if code == jwt_school_id and jwt_school_id != RESERVED_DEMO:
            mapped = code
        else:
            raise AppError(
                "VALIDATION",
                f"Unknown school_code '{code}'. Expected PRANAY for Pranay School Pune.",
                400,
                {"school_code": code, "allowed": list(SCHOOL_CODE_MAP), "jwt": jwt_school_id},
            )

    if mapped == RESERVED_DEMO or jwt_school_id == RESERVED_DEMO:
        raise AppError(
            "FORBIDDEN",
            "Cannot import into SCH-DEMO-01 via this path",
            403,
            {"school_code": code, "jwt": jwt_school_id},
        )

    if jwt_school_id != mapped:
        raise AppError(
            "VALIDATION",
            f"school_code maps to {mapped} but JWT schoolId is {jwt_school_id}",
            400,
            {"school_code": code, "expected": mapped, "jwt": jwt_school_id},
        )
    return mapped



def _validate_and_plan(rows: list[dict[str, str]], school_id: str) -> tuple[dict, list[dict], dict]:
    headers = set()
    if rows:
        headers = {k for k in rows[0].keys() if k != "_row"}
    forbidden = sorted(headers & FORBIDDEN_COLS)
    if forbidden:
        raise AppError(
            "VALIDATION",
            f"Forbidden column(s): {', '.join(forbidden)}",
            400,
            {"forbidden": forbidden},
        )
    missing = [c for c in SOT_REQUIRED if c not in headers]
    if missing:
        raise AppError(
            "VALIDATION",
            f"Missing required column(s): {', '.join(missing)}",
            400,
            {"missing": missing, "required": list(SOT_REQUIRED)},
        )

    errors: list[dict] = []
    # group by student_external_id for custody conflict detection
    by_student: dict[str, list[dict]] = {}
    planned_people: list[dict] = []
    student_meta: dict[str, dict] = {}

    for row in rows:
        row_no = int(row.get("_row") or 0)
        # F6 / AC-IMP-3 hard check first — never skip sparse court_order rows as "blank"
        custody_early = (row.get("custody_flag") or "").strip().lower().replace(" ", "_").replace("-", "_")
        custody_early = {"courtorder": "court_order", "court_order": "court_order", "court": "court_order"}.get(
            custody_early, custody_early
        ) if custody_early else ""
        gate_early = (row.get("gate_instruction") or "").strip()
        gate_early_blank = gate_early.lower() in {"", "-", "n/a", "na", "null", "none", "nil"}
        if custody_early == "court_order" and gate_early_blank:
            errors.append(
                _err(
                    row_no,
                    "gate_instruction required when custody_flag=court_order (F6 / AC-IMP-3)",
                    "gate_instruction",
                    "F6",
                )
            )
            # F6 rows must never be planned/committed
            continue
        if not any((row.get(k) or "").strip() for k in SOT_REQUIRED if k != "school_code"):
            continue
        ext = (row.get("student_external_id") or "").strip()
        name = (row.get("student_name") or "").strip()
        klass = (row.get("class") or "").strip()
        section = (row.get("section") or "").strip()
        person_name = (row.get("person_name") or "").strip()
        relation = (row.get("relation") or "").strip().lower()
        mobile_raw = (row.get("mobile") or "").strip()
        consent_version = (row.get("consent_version") or "").strip() or PICKUP_CONSENT_VERSION
        consent_at = (row.get("consent_at") or "").strip()
        row_errors: list[dict] = []
        if not ext:
            row_errors.append(_err(row_no, "student_external_id is required", "student_external_id"))
        if not name:
            row_errors.append(_err(row_no, "student_name is required", "student_name"))
        if not klass:
            row_errors.append(_err(row_no, "class is required", "class"))
        if not section:
            row_errors.append(_err(row_no, "section is required", "section"))
        if not person_name:
            row_errors.append(_err(row_no, "person_name is required", "person_name"))
        if relation not in RELATION_VALUES:
            row_errors.append(_err(row_no, f"relation must be one of {sorted(RELATION_VALUES)}", "relation"))
        try:
            mobile = normalize_mobile(mobile_raw) if mobile_raw else ""
        except Exception:
            mobile = ""
        if not mobile:
            row_errors.append(_err(row_no, "mobile is required", "mobile", "VALIDATION"))
        if not consent_at:
            row_errors.append(_err(row_no, "consent_at is required", "consent_at"))
        custody_raw = (row.get("custody_flag") or "").strip().lower()
        custody_raw = custody_raw.replace(" ", "_").replace("-", "_")
        custody = {"courtorder": "court_order", "court_order": "court_order", "court": "court_order"}.get(custody_raw, custody_raw) if custody_raw else "none"
        if custody not in CUSTODY_VALUES:
            row_errors.append(_err(row_no, "custody_flag must be none|restricted|court_order", "custody_flag"))
        gate_instruction = (row.get("gate_instruction") or "").strip()
        gate_blank = gate_instruction.lower() in {"", "-", "n/a", "na", "null", "none", "nil"}
        if custody == "court_order" and gate_blank:
            if not any(e.get("row") == row_no and e.get("code") == "F6" for e in errors):
                row_errors.append(
                    _err(
                        row_no,
                        "gate_instruction required when custody_flag=court_order (F6 / AC-IMP-3)",
                        "gate_instruction",
                        "F6",
                    )
                )
        if len(gate_instruction) > 280:
            row_errors.append(_err(row_no, "gate_instruction must be ≤280 chars", "gate_instruction"))
        try:
            person_active = _parse_bool(row.get("person_active", ""), True, "person_active")
            legal_hold = _parse_bool(row.get("legal_hold", ""), False, "legal_hold")
        except ValueError as e:
            row_errors.append(_err(row_no, str(e), "person_active"))
            person_active, legal_hold = True, False
        id_type = (row.get("id_type") or "").strip() or None
        if id_type and id_type not in ID_TYPE_VALUES:
            row_errors.append(_err(row_no, f"id_type must be one of {sorted(ID_TYPE_VALUES)}", "id_type"))
        id_last4 = (row.get("id_last4") or "").strip()
        if id_last4 and (not id_last4.isdigit() or len(id_last4) != 4):
            row_errors.append(_err(row_no, "id_last4 must be 4 digits", "id_last4"))

        if row_errors:
            errors.extend(row_errors)
            continue

        by_student.setdefault(ext, []).append(row)
        student_meta[ext] = {
            "student_external_id": ext,
            "name": name,
            "class": klass,
            "section": section,
            "legal_hold": legal_hold,
        }
        planned_people.append(
            {
                "row": row_no,
                "student_external_id": ext,
                "person_name": person_name,
                "relation": relation,
                "mobile": mobile,
                "mobile_raw": mobile_raw,
                "id_type": id_type,
                "id_last4": id_last4,
                "effective_from": (row.get("effective_from") or "").strip() or None,
                "effective_to": (row.get("effective_to") or "").strip() or None,
                "consent_version": consent_version,
                "consent_at": consent_at,
                "person_active": person_active,
                "custody_flag": custody,
                "gate_instruction": gate_instruction,
                "allowed_person_mobiles": (row.get("allowed_person_mobiles") or "").strip(),
                "blocked_person_mobiles": (row.get("blocked_person_mobiles") or "").strip(),
            }
        )

    # custody conflicts per student
    for ext, srows in by_student.items():
        flags = {(r.get("custody_flag") or "").strip().lower() or "none" for r in srows}
        instr = {(r.get("gate_instruction") or "").strip() for r in srows}
        if len(flags) > 1 or len(instr) > 1:
            # find first conflicting row after the first
            for r in srows[1:]:
                if ((r.get("custody_flag") or "").strip().lower() or "none") != next(iter(flags)) or (
                    (r.get("gate_instruction") or "").strip() != next(iter(instr))
                ):
                    errors.append(
                        _err(
                            int(r.get("_row") or 0),
                            f"Conflicting custody_flag/gate_instruction for student {ext}",
                            "custody_flag",
                            "CUSTODY_CONFLICT",
                        )
                    )

    summary = {
        "students": len(student_meta),
        "people": len(planned_people),
        "errorRows": len({e["row"] for e in errors}),
    }
    plan = {"students": student_meta, "people": planned_people}
    return plan, errors, summary


def _apply_plan(plan: dict, school_id: str, user_id: str) -> tuple[int, int]:
    ts = now_iso()
    imported = 0
    updated = 0
    student_ids: dict[str, str] = {}
    for ext, meta in plan["students"].items():
        existing = store.get_student_by_roster_id(school_id, ext)
        if existing:
            existing.update(
                {
                    "name": meta["name"],
                    "class": meta["class"],
                    "section": meta["section"],
                    "legalHold": meta["legal_hold"],
                    "active": True,
                    "updatedAt": ts,
                }
            )
            store.put_student(existing)
            sid = existing["id"]
            updated += 1
        else:
            sid = f"STU-{store.next_seq('student_seq'):04d}"
            store.put_student(
                {
                    "id": sid,
                    "schoolId": school_id,
                    "studentId": ext,
                    "name": meta["name"],
                    "class": meta["class"],
                    "section": meta["section"],
                    "active": True,
                    "enrollmentEndedAt": None,
                    "legalHold": meta["legal_hold"],
                    "createdAt": ts,
                    "updatedAt": ts,
                    "importedByUserId": user_id,
                }
            )
            if not store.get_custody_flag(sid):
                store.put_custody_flag(default_custody_flag(sid, school_id))
            imported += 1
        student_ids[ext] = sid

    people_created = 0
    people_updated = 0
    for p in plan["people"]:
        sid = student_ids[p["student_external_id"]]
        existing = store.get_authorized_person_by_mobile(school_id, sid, p["mobile"])
        payload = {
            "studentId": sid,
            "schoolId": school_id,
            "name": p["person_name"],
            "relation": p["relation"],
            "mobile": p["mobile"],
            "idType": p["id_type"],
            "idNumber": None,
            "idLast4": p["id_last4"] or None,
            "photoRef": None,
            "active": p["person_active"],
            "effectiveFrom": p["effective_from"],
            "effectiveTo": p["effective_to"],
            "pickupConsentVersion": p["consent_version"],
            "pickupConsentAt": p["consent_at"],
            "updatedAt": ts,
        }
        if existing:
            existing.update(payload)
            store.put_authorized_person(existing)
            people_updated += 1
        else:
            pid = f"AP-{store.next_seq('authorized_seq'):04d}"
            store.put_authorized_person({"id": pid, **payload, "createdAt": ts})
            people_created += 1

        # custody last-wins per student (already conflict-checked)
        flag = p["custody_flag"]
        gate_instruction = p["gate_instruction"]
        # resolve allow/block mobiles to person ids after all people upserted — second pass below
        student_ids[f"_custody_{p['student_external_id']}"] = sid  # keep
        student_ids[f"_flag_{p['student_external_id']}"] = flag  # type: ignore
        student_ids[f"_gate_{p['student_external_id']}"] = gate_instruction  # type: ignore
        student_ids[f"_allow_{p['student_external_id']}"] = p["allowed_person_mobiles"]  # type: ignore
        student_ids[f"_block_{p['student_external_id']}"] = p["blocked_person_mobiles"]  # type: ignore

    # write custody once per student
    for ext in plan["students"]:
        sid = student_ids[ext]
        flag = student_ids.get(f"_flag_{ext}") or "none"
        gate_instruction = student_ids.get(f"_gate_{ext}") or ""
        allow_raw = student_ids.get(f"_allow_{ext}") or ""
        block_raw = student_ids.get(f"_block_{ext}") or ""
        people = store.list_authorized_people(school_id, sid)
        by_mobile = {normalize_mobile(p["mobile"]): p["id"] for p in people if p.get("mobile")}

        def _ids(raw: str) -> list[str]:
            out = []
            for part in (raw or "").split(";"):
                m = normalize_mobile(part.strip()) if part.strip() else ""
                if m and m in by_mobile and by_mobile[m] not in out:
                    out.append(by_mobile[m])
            return out

        store.put_custody_flag(
            {
                "studentId": sid,
                "schoolId": school_id,
                "flag": flag,
                "gateInstruction": gate_instruction,
                "blockedPersonIds": _ids(block_raw),
                "allowedPersonIds": _ids(allow_raw) if allow_raw.strip() else None,
                "updatedByUserId": user_id,
                "updatedAt": ts,
            }
        )
        denormalize_blocked_by_custody(sid, school_id)

    # audit
    store.add_outbox(
        {
            "id": f"IMP-{store.next_seq('import_seq'):05d}",
            "schoolId": school_id,
            "event": "roster.import",
            "importedByUserId": user_id,
            "studentsCreated": imported,
            "studentsUpdated": updated,
            "peopleCreated": people_created,
            "peopleUpdated": people_updated,
            "status": "committed",
            "createdAt": ts,
            "meta": {"watermark": WATERMARK},
        }
    )
    return imported + people_created, updated + people_updated


def run_import(
    data: bytes,
    *,
    school_id: str,
    user_id: str,
    filename: Optional[str] = None,
    content_type: Optional[str] = None,
    dry_run: bool = True,
) -> dict:
    rows = _read_rows(data, filename, content_type)
    target = resolve_school(jwt_school_id=school_id, rows=rows)
    plan, errors, summary = _validate_and_plan(rows, target)

    imported = 0
    updated = 0
    if not dry_run:
        bad_rows = {e["row"] for e in errors}
        conflict_students = set()
        for e in errors:
            if e.get("code") == "CUSTODY_CONFLICT":
                for p in plan["people"]:
                    if p["row"] == e["row"]:
                        conflict_students.add(p["student_external_id"])
        bad_students = {
            p["student_external_id"]
            for p in plan["people"]
            if p["row"] in bad_rows
        } | conflict_students
        filtered = {
            "students": {k: v for k, v in plan["students"].items() if k not in bad_students},
            "people": [p for p in plan["people"] if p["student_external_id"] not in bad_students],
        }
        if filtered["students"]:
            imported, updated = _apply_plan(filtered, target, user_id)

    failed = summary["errorRows"]
    return {
        "schoolId": target,
        "school_code": next((k for k, v in SCHOOL_CODE_MAP.items() if v == target), target),
        "dryRun": dry_run,
        "ok": failed == 0,
        "imported": imported if not dry_run else 0,
        "updated": updated if not dry_run else 0,
        "failed": failed,
        "studentsInFile": summary["students"],
        "peopleInFile": summary["people"],
        "errors": errors,
        "meta": {"watermark": WATERMARK},
    }
